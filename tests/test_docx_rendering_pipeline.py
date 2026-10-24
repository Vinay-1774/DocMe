import io
import os
import sys
import tempfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import URL, create_engine, select, text
from sqlalchemy.orm import sessionmaker

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.app import config
from backend.app.auth import create_access_token, hash_password
from backend.app.database import Base, get_db
from backend.app.ingestion import (
    extract_docx_document_text,
    extract_text_and_metadata,
    chunk_document_sections,
)
from backend.app.main import app
from backend.app.models import Document, DocumentChunk, User

TEST_DB_NAME = "research_platform_test"

test_db_url = URL.create(
    drivername="postgresql+psycopg",
    username=config.DB_USER,
    password=config.DB_PASSWORD,
    host=config.DB_HOST,
    port=config.DB_PORT,
    database=TEST_DB_NAME,
)

test_engine = create_engine(test_db_url, pool_pre_ping=True)
TestSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def override_get_db():
    db = TestSessionLocal()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_get_db


@pytest.fixture(scope="session", autouse=True)
def setup_test_db():
    with test_engine.connect() as conn:
        conn.execute(text("CREATE EXTENSION IF NOT EXISTS vector;"))
        conn.commit()
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)


@pytest.fixture(autouse=True)
def clean_data():
    with TestSessionLocal() as session:
        session.query(DocumentChunk).delete()
        session.query(Document).delete()
        session.query(User).delete()
        session.commit()


@pytest.fixture
def test_docx_file():
    """Creates a temporary DOCX file with headings, paragraphs, bold runs, and a table."""
    doc = DocxDocument()
    doc.add_heading("Executive Summary", level=1)

    p1 = doc.add_paragraph("This preprint investigates ")
    bold_run = p1.add_run("quantum supremacy in optimization")
    bold_run.bold = True
    italic_run = p1.add_run(" and fault-tolerant architectures.")
    italic_run.italic = True

    doc.add_heading("Benchmark Results", level=2)

    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Baseline (ms)"
    table.cell(0, 2).text = "Quantum (ms)"

    table.cell(1, 0).text = "QAOA-128"
    table.cell(1, 1).text = "450"
    table.cell(1, 2).text = "32"

    table.cell(2, 0).text = "VQE-256"
    table.cell(2, 1).text = "920"
    table.cell(2, 2).text = "64"

    doc.add_paragraph("Conclusion: The benchmark confirms significant speedup across all tested workloads.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        doc.save(tmp.name)
        tmp_path = tmp.name

    yield tmp_path

    if os.path.isfile(tmp_path):
        os.unlink(tmp_path)


def test_docx_extraction_document_order_and_no_xml(test_docx_file):
    """
    Test 1 & 2:
    - Extracts heading 1, paragraph with bold text, heading 2, table rows, and conclusion in correct reading order.
    - Ensures no raw Word XML tags (<w:left>, <w:rPr>, <w:color>, <w:t>, etc.) are included in the extracted text.
    """
    extracted_text = extract_docx_document_text(test_docx_file)

    # 1. Verify correct document order
    pos_h1 = extracted_text.find("# Executive Summary")
    pos_p1 = extracted_text.find("quantum supremacy in optimization")
    pos_h2 = extracted_text.find("## Benchmark Results")
    pos_tbl_hdr = extracted_text.find("Model | Baseline (ms) | Quantum (ms)")
    pos_tbl_row1 = extracted_text.find("QAOA-128 | 450 | 32")
    pos_tbl_row2 = extracted_text.find("VQE-256 | 920 | 64")
    pos_concl = extracted_text.find("Conclusion: The benchmark confirms significant speedup")

    assert pos_h1 != -1, "Heading 1 should be extracted"
    assert pos_p1 != -1, "Paragraph with bold text should be extracted"
    assert pos_h2 != -1, "Heading 2 should be extracted"
    assert pos_tbl_hdr != -1, "Table header row should be extracted"
    assert pos_tbl_row1 != -1, "Table row 1 should be extracted"
    assert pos_tbl_row2 != -1, "Table row 2 should be extracted"
    assert pos_concl != -1, "Conclusion paragraph should be extracted"

    # Strictly verify reading order
    assert pos_h1 < pos_p1 < pos_h2 < pos_tbl_hdr < pos_tbl_row1 < pos_tbl_row2 < pos_concl

    # 2. Verify complete absence of raw Word XML tags
    assert "<w:" not in extracted_text, "No <w: tags should be present in extracted text"
    assert "</w:" not in extracted_text, "No </w: tags should be present in extracted text"
    assert "<w:rPr>" not in extracted_text, "No <w:rPr> should be present"
    assert "<w:left" not in extracted_text, "No <w:left> should be present"
    assert "<w:color" not in extracted_text, "No <w:color> should be present"
    assert "<w:t>" not in extracted_text, "No <w:t> tags should be present"


def test_docx_chunking_and_clean_text(test_docx_file):
    """
    Test 3: Verify chunk_document_sections generates clean text chunks
    without any XML leakage.
    """
    sections = extract_text_and_metadata(test_docx_file)
    assert len(sections) == 1
    assert sections[0]["page_number"] is None

    chunks = chunk_document_sections(sections, chunk_size=500, chunk_overlap=50)
    assert len(chunks) >= 1

    for chunk in chunks:
        content = chunk["content"]
        assert "<w:" not in content
        assert "</w:" not in content
        assert "<w:t" not in content
        assert "Executive Summary" in content or "Benchmark Results" in content or "Conclusion" in content


def test_authenticated_download_and_private_docx_access(test_docx_file):
    """
    Test 4: Authenticated downloads, permission checks (401, 403, 200) for private and public DOCX files.
    """
    client = TestClient(app)

    with TestSessionLocal() as session:
        user_owner = User(
            username="owner_user",
            email="owner@test.com",
            hashed_password=hash_password("password123"),
            is_verified=True,
        )
        user_other = User(
            username="other_user",
            email="other@test.com",
            hashed_password=hash_password("password123"),
            is_verified=True,
        )
        session.add_all([user_owner, user_other])
        session.commit()
        session.refresh(user_owner)
        session.refresh(user_other)

        owner_token = create_access_token(subject=user_owner.id)
        other_token = create_access_token(subject=user_other.id)

        # Create private document
        private_doc = Document(
            user_id=user_owner.id,
            title="Private Quantum Paper",
            file_path=test_docx_file,
            file_type="docx",
            visibility="private",
        )
        # Create public document
        public_doc = Document(
            user_id=user_owner.id,
            title="Public Quantum Paper",
            file_path=test_docx_file,
            file_type="docx",
            visibility="public",
        )
        session.add_all([private_doc, public_doc])
        session.commit()
        session.refresh(private_doc)
        session.refresh(public_doc)

        priv_id = private_doc.id
        pub_id = public_doc.id

    # 1. Unauthenticated download of private document -> 401
    res_unauth = client.get(f"/documents/{priv_id}/download")
    assert res_unauth.status_code == 401, "Unauthenticated access to private doc must return 401"

    # 2. Unauthorized other user download of private document -> 403
    res_forbidden = client.get(
        f"/documents/{priv_id}/download",
        headers={"Authorization": f"Bearer {other_token}"},
    )
    assert res_forbidden.status_code == 403, "Other user access to private doc must return 403"

    # 3. Owner download of private document -> 200
    res_owner = client.get(
        f"/documents/{priv_id}/download?inline=true",
        headers={"Authorization": f"Bearer {owner_token}"},
    )
    assert res_owner.status_code == 200, "Owner must be able to download private document"
    assert len(res_owner.content) > 0

    # 4. Other user download of public document -> 200
    res_pub = client.get(
        f"/documents/{pub_id}/download?inline=true",
        headers={"Authorization": f"Bearer {other_token}"},
    )
    assert res_pub.status_code == 200, "Any user can download public document"
    assert len(res_pub.content) > 0


def test_ai_chat_chunks_clean_text_retrieval(test_docx_file):
    """
    Test 6: AI Chat retrieval uses clean DOCX chunks without raw XML tags.
    """
    sections = extract_text_and_metadata(test_docx_file)
    chunks = chunk_document_sections(sections, chunk_size=300, chunk_overlap=50)

    with TestSessionLocal() as session:
        user = User(
            username="rag_user",
            email="rag@test.com",
            hashed_password=hash_password("password123"),
            is_verified=True,
        )
        session.add(user)
        session.commit()
        session.refresh(user)

        doc = Document(
            user_id=user.id,
            title="Clean RAG DOCX",
            file_path=test_docx_file,
            file_type="docx",
            visibility="public",
        )
        session.add(doc)
        session.commit()
        session.refresh(doc)

        for chunk_data in chunks:
            chunk_obj = DocumentChunk(
                document_id=doc.id,
                user_id=user.id,
                content=chunk_data["content"],
                chunk_index=chunk_data["chunk_index"],
                page_number=None,
                embedding=[0.01] * 768,
            )
            session.add(chunk_obj)
        session.commit()

        # Retrieve stored chunks from database
        stored_chunks = session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == doc.id)
        ).all()

        assert len(stored_chunks) > 0
        combined_content = " ".join(sc.content for sc in stored_chunks)
        assert "<w:" not in combined_content
        assert "</w:" not in combined_content
        assert "<w:rPr>" not in combined_content
        assert "<w:t>" not in combined_content
        for expected_word in ["Executive", "quantum", "QAOA-128", "Benchmark", "speedup"]:
            assert expected_word in combined_content, f"Expected {expected_word} to be present in stored chunks"
